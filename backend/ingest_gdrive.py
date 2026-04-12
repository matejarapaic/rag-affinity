"""
Google Drive → RAG Ingestion Pipeline
======================================
Downloads PDFs, Word docs, Google Docs, and spreadsheets from a Google
Drive folder, chunks them, embeds with fastembed, and upserts into Qdrant.

Usage:
    python ingest_gdrive.py                  # ingest all new/changed files
    python ingest_gdrive.py --force          # re-ingest everything
    python ingest_gdrive.py --file-id <id>   # ingest a single file
"""

import os
import io
import json
import hashlib
import argparse
import logging
from pathlib import Path
from datetime import datetime

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import fitz  # PyMuPDF
import docx  # python-docx
import openpyxl
import csv
from fastembed import TextEmbedding
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance, VectorParams, PointStruct, Filter,
    FieldCondition, MatchValue
)

# ---------------------------------------------------------------------------
# Configuration — override via environment variables or edit directly
# ---------------------------------------------------------------------------

import sys as _sys

def _resolve_service_account():
    """Find service_account.json whether running normally or frozen by PyInstaller."""
    env_val = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    if env_val:
        return env_val
    # When frozen, data files land in sys._MEIPASS
    if getattr(_sys, 'frozen', False):
        return os.path.join(_sys._MEIPASS, "service_account.json")
    # Dev: same directory as this file
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "service_account.json")

GDRIVE_FOLDER_ID   = os.getenv("GDRIVE_FOLDER_ID", "1-lekJE_VDOpnp13OUXtCXZc4Nx71xGFb")
SERVICE_ACCOUNT_FILE = _resolve_service_account()

QDRANT_URL         = os.getenv("QDRANT_URL", "http://localhost:6333")
QDRANT_API_KEY     = os.getenv("QDRANT_API_KEY", "")           # leave blank if local
COLLECTION_NAME    = os.getenv("QDRANT_COLLECTION", "affinity_rag")

EMBED_MODEL        = "BAAI/bge-small-en-v1.5"                  # fast, good quality
CHUNK_SIZE         = 512    # characters per chunk
CHUNK_OVERLAP      = 64     # overlap between chunks
BATCH_SIZE         = 32     # embeddings per batch

STATE_FILE         = ".ingest_state.json"  # tracks file hashes to avoid re-ingestion

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Google Drive helpers
# ---------------------------------------------------------------------------

def get_drive_service():
    """Authenticate with Google Drive using a service account."""
    scopes = ["https://www.googleapis.com/auth/drive.readonly"]
    creds = service_account.Credentials.from_service_account_file(
        SERVICE_ACCOUNT_FILE, scopes=scopes
    )
    return build("drive", "v3", credentials=creds)


# MIME types we support and how to handle them
SUPPORTED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",  # .docx
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",        # .xlsx
    "application/vnd.google-apps.document",      # Google Doc
    "application/vnd.google-apps.spreadsheet",   # Google Sheet
    "application/vnd.google-apps.presentation",  # Google Slides
    "text/plain",
}

# Google Workspace files must be exported, not downloaded directly
GOOGLE_EXPORT_MIME = {
    "application/vnd.google-apps.document":     "text/plain",
    "application/vnd.google-apps.spreadsheet":  "text/csv",
    "application/vnd.google-apps.presentation": "text/plain",
}


def list_files(service, folder_id: str) -> list[dict]:
    """List all supported files in a Drive folder."""
    files = []
    mime_filter = " or ".join(f"mimeType='{m}'" for m in SUPPORTED_MIME_TYPES)
    query = f"'{folder_id}' in parents and ({mime_filter}) and trashed=false"
    page_token = None
    while True:
        resp = service.files().list(
            q=query,
            fields="nextPageToken, files(id, name, mimeType, modifiedTime, md5Checksum)",
            pageToken=page_token,
            pageSize=100,
        ).execute()
        files.extend(resp.get("files", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    log.info(f"Found {len(files)} supported files in Drive folder")
    return files


def download_file(service, file_id: str, mime_type: str) -> bytes:
    """Download or export a file from Drive into memory."""
    export_mime = GOOGLE_EXPORT_MIME.get(mime_type)
    if export_mime:
        # Google Workspace files must be exported
        request = service.files().export_media(fileId=file_id, mimeType=export_mime)
    else:
        request = service.files().get_media(fileId=file_id)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Text extraction — one function per file type
# ---------------------------------------------------------------------------

def extract_text_pdf(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    pages = [page.get_text("text").strip() for page in doc if page.get_text("text").strip()]
    doc.close()
    return "\n\n".join(pages)


def extract_text_docx(data: bytes) -> str:
    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join(str(c) if c is not None else "" for c in row)
            if line.strip():
                rows.append(line)
    return "\n".join(rows)


def extract_text_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join("\t".join(row) for row in reader if any(row))


def extract_text_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def extract_text(data: bytes, mime_type: str) -> str:
    """Dispatch to the correct extractor based on MIME type."""
    # Google Workspace files are exported as text/plain or text/csv
    if mime_type in ("application/vnd.google-apps.document", "application/vnd.google-apps.presentation"):
        return extract_text_plain(data)
    if mime_type == "application/vnd.google-apps.spreadsheet":
        return extract_text_csv(data)
    if mime_type == "application/pdf":
        return extract_text_pdf(data)
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_docx(data)
    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_text_xlsx(data)
    if mime_type == "text/plain":
        return extract_text_plain(data)
    raise ValueError(f"Unsupported MIME type: {mime_type}")


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


# ---------------------------------------------------------------------------
# Vector DB helpers
# ---------------------------------------------------------------------------

def get_qdrant_client() -> QdrantClient:
    if QDRANT_API_KEY:
        return QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)
    return QdrantClient(url=QDRANT_URL)


def ensure_collection(client: QdrantClient, vector_size: int):
    """Create the Qdrant collection if it doesn't exist."""
    existing = [c.name for c in client.get_collections().collections]
    if COLLECTION_NAME not in existing:
        client.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=VectorParams(size=vector_size, distance=Distance.COSINE),
        )
        log.info(f"Created Qdrant collection: {COLLECTION_NAME}")
    else:
        log.info(f"Using existing Qdrant collection: {COLLECTION_NAME}")


def delete_file_chunks(client: QdrantClient, file_id: str):
    """Remove all existing chunks for a file before re-ingesting."""
    client.delete(
        collection_name=COLLECTION_NAME,
        points_selector=Filter(
            must=[FieldCondition(key="file_id", match=MatchValue(value=file_id))]
        ),
    )


def upsert_chunks(client: QdrantClient, embedder, chunks: list[str], metadata: dict):
    """Embed chunks in batches and upsert into Qdrant."""
    points = []
    chunk_id_base = int(hashlib.md5(metadata["file_id"].encode()).hexdigest(), 16) % (10**12)

    for i in range(0, len(chunks), BATCH_SIZE):
        batch = chunks[i : i + BATCH_SIZE]
        embeddings = list(embedder.embed(batch))

        for j, (chunk, vector) in enumerate(zip(batch, embeddings)):
            point_id = chunk_id_base + i + j
            points.append(
                PointStruct(
                    id=point_id,
                    vector=vector.tolist(),
                    payload={
                        **metadata,
                        "chunk_id":    point_id,
                        "chunk_index": i + j,
                        "text": chunk,
                        "source": "google_drive",
                        "ingested_at": datetime.utcnow().isoformat(),
                    },
                )
            )

    client.upsert(collection_name=COLLECTION_NAME, points=points)
    log.info(f"  Upserted {len(points)} chunks for '{metadata['file_name']}'")


# ---------------------------------------------------------------------------
# State tracking (skip unchanged files)
# ---------------------------------------------------------------------------

def load_state() -> dict:
    if Path(STATE_FILE).exists():
        with open(STATE_FILE) as f:
            return json.load(f)
    return {}


def save_state(state: dict):
    with open(STATE_FILE, "w") as f:
        json.dump(state, f, indent=2)


# ---------------------------------------------------------------------------
# Main ingestion logic
# ---------------------------------------------------------------------------

def ingest_file(service, client, embedder, file_meta: dict, force: bool = False):
    """Download, parse, embed, and store a single Drive file."""
    file_id   = file_meta["id"]
    file_name = file_meta["name"]
    mime_type = file_meta.get("mimeType", "")
    # Google Workspace files don't have an md5Checksum — use modifiedTime instead
    checksum  = file_meta.get("md5Checksum") or file_meta.get("modifiedTime", "")

    state = load_state()

    # Migrate old string-format state entries to dict format
    raw = state.get(file_id, {})
    entry = raw if isinstance(raw, dict) else {"checksum": raw, "name": ""}

    # Skip if checksum AND name are both unchanged
    if not force and entry.get("checksum") == checksum and entry.get("name") == file_name:
        log.info(f"  Skipping (unchanged): {file_name}")
        return
    if entry.get("name") and entry.get("name") != file_name:
        log.info(f"  Name changed: '{entry['name']}' → '{file_name}', re-ingesting")

    log.info(f"  Ingesting [{mime_type}]: {file_name}")

    # Download / export
    file_bytes = download_file(service, file_id, mime_type)

    # Parse
    try:
        text = extract_text(file_bytes, mime_type)
    except Exception as e:
        log.warning(f"  Failed to extract text from {file_name}: {e} — skipping")
        return

    if not text.strip():
        log.warning(f"  No text extracted from {file_name} — skipping")
        return

    # Chunk
    chunks = chunk_text(text)
    log.info(f"  Split into {len(chunks)} chunks")

    # Remove old vectors for this file
    delete_file_chunks(client, file_id)

    # Embed + store
    metadata = {
        "file_id":       file_id,
        "file_name":     file_name,
        "modified_time": file_meta.get("modifiedTime", ""),
        "checksum":      checksum,
        "doc_id":        file_id,
        "author_name":   "",
        "matter_id":     "",
        "doc_type":      "unknown",
    }
    upsert_chunks(client, embedder, chunks, metadata)

    # Update state
    state[file_id] = {"checksum": checksum, "name": file_name}
    save_state(state)


def run(folder_id: str = None, force: bool = False, single_file_id: str = None):
    folder_id = folder_id or GDRIVE_FOLDER_ID
    if not folder_id and not single_file_id:
        raise ValueError("Set GDRIVE_FOLDER_ID env var or pass --folder-id")

    log.info("Initialising Google Drive service...")
    service = get_drive_service()

    log.info("Initialising fastembed model...")
    embedder = TextEmbedding(model_name=EMBED_MODEL)

    log.info("Connecting to Qdrant...")
    client = get_qdrant_client()

    # Determine vector size from a test embed
    sample = list(embedder.embed(["test"]))[0]
    ensure_collection(client, len(sample))

    if single_file_id:
        meta = service.files().get(
            fileId=single_file_id,
            fields="id, name, mimeType, modifiedTime, md5Checksum"
        ).execute()
        ingest_file(service, client, embedder, meta, force=True)
    else:
        files = list_files(service, folder_id)
        for file_meta in files:
            ingest_file(service, client, embedder, file_meta, force=force)

    log.info("✓ Ingestion complete")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest Google Drive PDFs into RAG")
    parser.add_argument("--folder-id",  help="Override GDRIVE_FOLDER_ID env var")
    parser.add_argument("--file-id",    help="Ingest a single file by Drive ID")
    parser.add_argument("--force",      action="store_true", help="Re-ingest all files")
    args = parser.parse_args()

    run(folder_id=args.folder_id, force=args.force, single_file_id=args.file_id)
